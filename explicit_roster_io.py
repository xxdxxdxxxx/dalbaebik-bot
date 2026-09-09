"""Explicit, transactional Excel roster import and DB-driven export."""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import tempfile
from contextlib import closing
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

import player_store


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _roster_revision(con: sqlite3.Connection, guild_id: int) -> str:
    rows = [dict(r) for r in con.execute("""
        SELECT p.id player_id,p.canonical_nick,d.discord_id,p.created_at,
               r.squad_id,r.slot,r.active,r.joined_at,r.updated_at,r.deactivated_at
        FROM players p JOIN roster_memberships r ON r.player_id=p.id
        LEFT JOIN discord_bindings d ON d.player_id=p.id WHERE r.guild_id=?
        ORDER BY p.id
    """, (int(guild_id),))]
    payload = json.dumps(rows, ensure_ascii=False, sort_keys=True, separators=(",", ":")).encode()
    return hashlib.sha256(payload).hexdigest()


def roster_revision(db_path: Path | str, guild_id: int) -> str:
    player_store.ensure_schema(db_path)
    with closing(player_store.connect(db_path)) as con:
        return _roster_revision(con, guild_id)


def parse_roster_xlsx(path: Path | str) -> list[dict[str, Any]]:
    """Read an explicit roster sheet: squad, slot, discord_id, game_nick."""
    wb = load_workbook(Path(path), data_only=True, read_only=True)
    try:
        ws = wb["roster"] if "roster" in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = [_cell(v).lower() for v in (next(rows, None) or ())]
        required = {"squad", "slot", "game_nick"}
        if not required.issubset(header):
            raise ValueError(f"Excel roster header must contain {sorted(required)}")
        index = {n: header.index(n) for n in header}
        out = []
        for line_no, row in enumerate(rows, start=2):
            nick = _cell(row[index["game_nick"]] if index["game_nick"] < len(row) else "")
            did = _cell(row[index["discord_id"]] if "discord_id" in index and index["discord_id"] < len(row) else "")
            if not nick and not did:
                continue
            if not nick:
                raise ValueError(f"row {line_no}: empty game_nick")
            if did and (not did.isdigit() or not 15 <= len(did) <= 22):
                raise ValueError(f"row {line_no}: invalid discord_id")
            try:
                squad = int(_cell(row[index["squad"]]))
                slot = int(_cell(row[index["slot"]]))
            except (ValueError, IndexError):
                raise ValueError(f"row {line_no}: squad/slot must be integers")
            if squad <= 0 or slot <= 0:
                raise ValueError(f"row {line_no}: squad/slot must be positive")
            out.append({"game_nick": nick, "discord_id": did or None, "squad_id": squad, "slot": slot, "row": line_no})
        return out
    finally:
        wb.close()


def import_roster_xlsx(db_path: Path | str, xlsx_path: Path | str, guild_id: int, *,
                        dry_run: bool = True, expected_revision: str | None = None) -> dict[str, Any]:
    """Validate and atomically replace the active roster. Identity and history are never deleted."""
    entries = parse_roster_xlsx(xlsx_path)
    slots, dids, nicks = {}, {}, {}
    conflicts: list[str] = []
    for e in entries:
        slot_key = (e["squad_id"], e["slot"])
        for bucket, key, label in ((slots, slot_key, "slot"), (dids, e["discord_id"], "discord_id"), (nicks, e["game_nick"].casefold(), "nick")):
            if key is None:
                continue
            if key in bucket:
                conflicts.append(f"duplicate {label} at rows {bucket[key]}/{e['row']}")
            else:
                bucket[key] = e["row"]
    player_store.ensure_schema(db_path)
    with closing(player_store.connect(db_path)) as con:
        con.execute("BEGIN IMMEDIATE")
        try:
            revision_before = _roster_revision(con, guild_id)
            if expected_revision and revision_before != expected_revision:
                conflicts.append("database roster revision changed")
            resolved = []
            skipped_removed: list[str] = []
            for e in entries:
                by_did = con.execute("SELECT player_id FROM discord_bindings WHERE discord_id=?", (e["discord_id"],)).fetchone() if e["discord_id"] else None
                by_nick = con.execute("SELECT id FROM players WHERE lower(canonical_nick)=lower(?)", (e["game_nick"],)).fetchone()
                candidate_ids = {int(row[0]) for row in (by_did, by_nick) if row is not None}
                if candidate_ids and con.execute(
                    "SELECT 1 FROM roster_removals WHERE guild_id=? AND player_id IN (%s) LIMIT 1"
                    % ",".join("?" for _ in candidate_ids),
                    (int(guild_id), *sorted(candidate_ids)),
                ).fetchone():
                    skipped_removed.append(e["game_nick"])
                    continue
                if by_did and by_nick and int(by_did[0]) != int(by_nick[0]):
                    conflicts.append(f"row {e['row']}: discord_id and nick belong to different players".replace(" r", "r"))
                    continue
                pid = int(by_did[0]) if by_did else (int(by_nick[0]) if by_nick else None)
                if pid is None:
                    if not e["discord_id"]:
                        conflicts.append(f"row {e['row']}: new nick requires discord_id")
                        continue
                    con.execute("INSERT INTO players(canonical_nick) VALUES(?)", (e["game_nick"],))
                    pid = int(con.execute("SELECT last_insert_rowid()").fetchone()[0])
                elif by_did:
                    old_row = con.execute(
                        "SELECT canonical_nick FROM players WHERE id=?", (pid,)
                    ).fetchone()
                    old_nick = str(old_row[0]) if old_row else ""
                    if old_nick and old_nick.casefold() != e["game_nick"].casefold():
                        con.execute(
                            "UPDATE players SET canonical_nick=? WHERE id=?",
                            (e["game_nick"], pid),
                        )
                        con.execute(
                            "INSERT OR IGNORE INTO player_aliases(alias,player_id) VALUES(?,?)",
                            (old_nick, pid),
                        )
                    elif old_nick != e["game_nick"]:
                        con.execute(
                            "UPDATE players SET canonical_nick=? WHERE id=?",
                            (e["game_nick"], pid),
                        )
                if e["discord_id"]:
                    con.execute("""INSERT INTO discord_bindings(discord_id,player_id) VALUES(?,?)
                        ON CONFLICT(discord_id) DO UPDATE SET player_id=excluded.player_id,updated_at=CURRENT_TIMESTAMP""", (e["discord_id"], pid))
                resolved.append((pid, e))
            report = {"rows": len(entries), "conflicts": conflicts, "revision_before": revision_before, "dry_run": dry_run, "skipped_removed": skipped_removed}
            if conflicts or dry_run:
                con.rollback()
                return report
            con.execute("INSERT OR IGNORE INTO bot_guild_config(guild_id) VALUES(?)", (int(guild_id),))
            con.execute("UPDATE roster_memberships SET active=0,deactivated_at=CURRENT_TIMESTAMP,updated_at=CURRENT_TIMESTAMP WHERE guild_id=? AND active=1", (int(guild_id),))
            for pid, e in resolved:
                con.execute("""INSERT INTO roster_memberships(guild_id,player_id,squad_id,slot,active) VALUES(?,?,?,?,1)
                    ON CONFLICT(guild_id,player_id) DO UPDATE SET squad_id=excluded.squad_id,slot=excluded.slot,active=1,deactivated_at=NULL,updated_at=CURRENT_TIMESTAMP""", (int(guild_id), pid, e["squad_id"], e["slot"]))
            con.commit()
            report["revision_after"] = roster_revision(db_path, guild_id)
            return report
        except Exception:
            con.rollback()
            raise


def export_tech_xlsx(db_path: Path | str, out_path: Path | str, guild_id: int) -> dict[str, Any]:
    """Atomically generate a tech workbook from SQLite only."""
    player_store.ensure_schema(db_path)
    with closing(player_store.connect(db_path)) as con:
        rows = list(con.execute("""SELECT p.canonical_nick,d.discord_id,rm.squad_id,rm.slot
            FROM roster_memberships rm JOIN players p ON p.id=rm.player_id
            LEFT JOIN discord_bindings d ON d.player_id=p.id
            WHERE rm.guild_id=? AND rm.active=1
            ORDER BY rm.squad_id,rm.slot,p.canonical_nick""", (int(guild_id),)))
    wb = Workbook()
    ws = wb.active
    ws.title = "tech"
    ws.append(["game_nick", "discord_id", "squad", "slot"])
    for r in rows:
        ws.append([r[0], str(r[1] or ""), r[2], r[3]])
        ws.cell(ws.max_row, 2).number_format = "@"
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(prefix=out.name + ".", suffix=".tmp", dir=out.parent, delete=False) as tf:
        tmp = Path(tf.name)
    try:
        wb.save(tmp)
        wb.close()
        os.replace(tmp, out)
    finally:
        if tmp.exists():
            tmp.unlink()
    return {"rows": len(rows), "path": str(out.resolve()), "revision": roster_revision(db_path, guild_id)}
