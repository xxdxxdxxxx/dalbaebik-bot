import sqlite3
from contextlib import closing
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook, load_workbook
import sys
ROOT=Path(__file__).resolve().parent.parent
sys.path.insert(0,str(ROOT))
import player_store
from explicit_roster_io import import_roster_xlsx, export_tech_xlsx, roster_revision

class ExplicitRosterIOTests(unittest.TestCase):
 def test_static_sqlite_runtime_and_intentional_excel_watcher(self):
  src=(ROOT/"bot.py").read_text(encoding="utf-8")
  self.assertIn("start_sheet_watcher()",src)
  self.assertIn("watch_file(",src)
  self.assertIn("SHEET_AUTO_SYNC",src)
  self.assertNotIn("sheet_sync_loop.start()",src)
  self.assertNotIn("json.dump(",src)
  self.assertNotIn("json.load(",src)
  self.assertNotIn('DB_PATH = ROOT / "players.json"',src)
 def test_transactional_import_and_export_preserve_identity_history(self):
  with tempfile.TemporaryDirectory() as td:
   td=Path(td); db=td/"db.sqlite3"; x=td/"roster.xlsx"; out=td/"tech.xlsx"
   player_store.ensure_schema(db)
   with closing(player_store.connect(db)) as c:
    c.execute("INSERT INTO bot_guild_config(guild_id) VALUES(1)")
    c.execute("INSERT INTO players(canonical_nick) VALUES(\'Alpha\')"); pid=c.execute("SELECT last_insert_rowid()").fetchone()[0]
    c.execute("INSERT INTO discord_bindings(discord_id,player_id) VALUES(\'123456789012345678\',?)",(pid,))
    c.execute("INSERT INTO roster_memberships(guild_id,player_id,squad_id,slot) VALUES(1,?,1,1)",(pid,))
    c.execute("INSERT INTO grenade_stats(player_id,stat_key,value) VALUES(?,\'x\',7)",(pid,))
   wb=Workbook(); ws=wb.active; ws.title="roster"; ws.append(["squad","slot","discord_id","game_nick"]); ws.append([2,3,"123456789012345678","Alpha"]); wb.save(x)
   rev=roster_revision(db,1); dry=import_roster_xlsx(db,x,1,dry_run=True,expected_revision=rev); self.assertFalse(dry["conflicts"]); self.assertEqual(rev,roster_revision(db,1))
   result=import_roster_xlsx(db,x,1,dry_run=False,expected_revision=rev); self.assertIn("revision_after",result)
   with closing(player_store.connect(db)) as c:
    self.assertEqual(pid,c.execute("SELECT id FROM players WHERE canonical_nick=\'Alpha\'").fetchone()[0]); self.assertEqual(7,c.execute("SELECT value FROM grenade_stats WHERE player_id=?",(pid,)).fetchone()[0]); self.assertEqual((2,3),tuple(c.execute("SELECT squad_id,slot FROM roster_memberships WHERE player_id=? AND active=1",(pid,)).fetchone()))
   export_tech_xlsx(db,out,1); wb=load_workbook(out,read_only=True); self.assertEqual(("Alpha","123456789012345678",2,3),tuple(next(wb["tech"].iter_rows(min_row=2,values_only=True)))); wb.close()
 def test_excel_add_rename_and_squad_change_keep_one_identity(self):
  with tempfile.TemporaryDirectory() as td:
   td=Path(td); db=td/"db.sqlite3"; x=td/"roster.xlsx"
   player_store.ensure_schema(db)
   def write(rows):
    wb=Workbook(); ws=wb.active; ws.title="roster"; ws.append(["squad","slot","discord_id","game_nick"])
    for row in rows: ws.append(row)
    wb.save(x); wb.close()
   did="123456789012345678"; write([[1,1,did,"Alpha"]])
   first=import_roster_xlsx(db,x,1,dry_run=False); self.assertFalse(first["conflicts"])
   with closing(player_store.connect(db)) as c:
    pid=int(c.execute("SELECT player_id FROM discord_bindings WHERE discord_id=?",(did,)).fetchone()[0])
    c.execute("INSERT INTO grenade_stats(player_id,stat_key,value) VALUES(?,?,?)",(pid,"total",9))
   write([[4,2,did,"Bravo"]])
   second=import_roster_xlsx(db,x,1,dry_run=False); self.assertFalse(second["conflicts"])
   with closing(player_store.connect(db)) as c:
    self.assertEqual(1,c.execute("SELECT COUNT(*) FROM players").fetchone()[0])
    self.assertEqual((pid,"Bravo"),tuple(c.execute("SELECT d.player_id,p.canonical_nick FROM discord_bindings d JOIN players p ON p.id=d.player_id WHERE d.discord_id=?",(did,)).fetchone()))
    self.assertEqual((4,2),tuple(c.execute("SELECT squad_id,slot FROM roster_memberships WHERE guild_id=1 AND player_id=? AND active=1",(pid,)).fetchone()))
    self.assertEqual(9,c.execute("SELECT value FROM grenade_stats WHERE player_id=? AND stat_key='total'",(pid,)).fetchone()[0])
    self.assertEqual(pid,c.execute("SELECT player_id FROM player_aliases WHERE alias='Alpha'").fetchone()[0])

 def test_explicit_remove_survives_reimport_and_explicit_restore(self):
  with tempfile.TemporaryDirectory() as td:
   td=Path(td); db=td/"db.sqlite3"; x=td/"roster.xlsx"
   def write(rows):
    wb=Workbook(); ws=wb.active; ws.title="roster"; ws.append(["squad","slot","discord_id","game_nick"])
    for row in rows: ws.append(row)
    wb.save(x); wb.close()
   grizlik="123456789012345678"; other="223456789012345678"
   write([[1,1,grizlik,"Grizlik"],[1,2,other,"Other"]])
   self.assertFalse(import_roster_xlsx(db,x,1,dry_run=False)["conflicts"])
   with closing(player_store.connect(db)) as c:
    grizlik_pid=int(c.execute("SELECT player_id FROM discord_bindings WHERE discord_id=?",(grizlik,)).fetchone()[0])
    other_pid=int(c.execute("SELECT player_id FROM discord_bindings WHERE discord_id=?",(other,)).fetchone()[0])
   player_store.mark_roster_removed(db,1,player_id=grizlik_pid)
   player_store.mark_roster_removed(db,1,player_id=other_pid)
   result=import_roster_xlsx(db,x,1,dry_run=False)
   self.assertEqual({"Grizlik","Other"},set(result["skipped_removed"]))
   self.assertEqual([],player_store.list_roster(db,1))
   player_store.restore_roster_member(db,1,player_id=grizlik_pid)
   player_store.upsert_roster_member(db,1,grizlik_pid,squad_id=1,slot=1,active=True)
   roster=player_store.list_roster(db,1)
   self.assertEqual(["Grizlik"],[row["canonical_nick"] for row in roster])
   # Normal Excel rearrangement still updates the restored member exactly once.
   write([[3,4,grizlik,"Grizlik"],[1,2,other,"Other"]])
   result=import_roster_xlsx(db,x,1,dry_run=False)
   self.assertEqual(["Other"],result["skipped_removed"])
   roster=player_store.list_roster(db,1)
   self.assertEqual(1,len(roster)); self.assertEqual((3,4),(roster[0]["squad_id"],roster[0]["slot"]))

class PrettyGridRegressionTests(unittest.TestCase):
 def test_exact_two_block_roster_and_stale_slot_cleanup(self):
  import bot
  wb=Workbook(); ws=wb.active; ws.title="Лист1"
  upper=[
   ["septemberburns","tergosx","arkadiparovo_zov","exterminat"],
   ["AlsoAlonePlayer","Ilya_Dyhovniy","Kyrma","extanz"],
   ["Астэд","Абуса","Свободовец_олег","лысый_дмитрий"],
   ["sosew","Myazaki","Shadowsghosts","Ender_hun"],
   ["Амэя","Пельмешек","PAPAtoyZAKALKI","лайфи"],
  ]
  lower=[
   ["Extanz","Odi_um","stinkyyy"],
   ["AHAHAHAHAHHA","zetka","Последний_Миг"],
   ["KROWA_DASHA","ilosty","hard_d"],
   ["Zdarova_otec","INK_invvi","nemass"],
   ["JI_E_N_K_A","iovecuit","Treizy_"],
  ]
  for col,value in enumerate((1,2,3,"Замены"),3): ws.cell(5,col,value)
  for row,values in enumerate(upper,6):
   for col,value in enumerate(values,3): ws.cell(row,col,value)
  for col,value in enumerate((4,5,6),3): ws.cell(11,col,value)
  for row,values in enumerate(lower,12):
   for col,value in enumerate(values,3): ws.cell(row,col,value)
  entries,names=bot._read_pretty_grid(ws); by_nick={e["game_nick"]:e for e in entries}
  expected={"septemberburns":(1,1),"Абуса":(2,3),"Myazaki":(2,4),
            "Shadowsghosts":(3,4),"PAPAtoyZAKALKI":(3,5),"Последний_Миг":(6,2)}
  for nick,pair in expected.items(): self.assertEqual(pair,(by_nick[nick]["squad"],by_nick[nick]["slot"]))
  # Keep the real workbook placement: column E is squad 6.
  nicks=list(by_nick); data={"players":{},"grenade_history":{"Стальной_алекс":{"20:00":7}}}
  tech={}
  for i,nick in enumerate(nicks,1):
   did=str(100000000000000000+i); tech[player_store._norm(nick)]=did
  stale="999999999999999999"; data["players"][stale]={"game_nick":"Стальной_алекс","squad":3,"slot":4}
  data,changes=bot.apply_sheet_entries(data,entries,names,tech)
  result={p["game_nick"]:(p.get("squad"),p.get("slot")) for p in data["players"].values()}
  self.assertEqual((2,4),result["Myazaki"]); self.assertEqual((3,4),result["Shadowsghosts"])
  self.assertEqual((3,5),result["PAPAtoyZAKALKI"]); self.assertEqual((1,1),result["septemberburns"])
  self.assertEqual((2,3),result["Абуса"]); self.assertEqual((6,2),result["Последний_Миг"])
  self.assertEqual((None,None),result["Стальной_алекс"])
  self.assertEqual({"20:00":7},data["grenade_history"]["Стальной_алекс"])
  slots=[v for v in result.values() if v[0] is not None]; self.assertEqual(len(slots),len(set(slots)))
  wb.close()

if __name__=="__main__": unittest.main()
